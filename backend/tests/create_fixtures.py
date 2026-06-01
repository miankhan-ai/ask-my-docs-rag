"""Generate the sample test fixtures (txt, md, docx, pdf). Run once."""
from pathlib import Path

import docx
from fpdf import FPDF

FIXTURES = Path(__file__).parent / "fixtures"
FIXTURES.mkdir(parents=True, exist_ok=True)

(FIXTURES / "sample.txt").write_text(
    "Hello from plain text document. It supports PDF, DOCX, Markdown, and plain text files.",
    encoding="utf-8",
)

(FIXTURES / "sample.md").write_text(
    "# Hello\n\nHello from markdown document.",
    encoding="utf-8",
)

doc = docx.Document()
doc.add_paragraph("Hello from docx document.")
doc.save(str(FIXTURES / "sample.docx"))

pdf = FPDF()
pdf.add_page()
pdf.set_font("Helvetica", size=12)
pdf.cell(200, 10, "Hello from PDF document.")
pdf.output(str(FIXTURES / "sample.pdf"))

print("Fixtures created in", FIXTURES)
